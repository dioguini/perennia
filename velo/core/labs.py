"""Análises clínicas na estrutura do teu Excel.

Duas folhas:

* «readings» — uma linha por resultado: tipo (Rotina…), laboratório, data,
  área (Hematologia, Bioquímica…), exame (Hemograma…), subexame
  (Eritrograma…), parâmetro e valor. O valor pode ser um número
  («11,1»), um número com a percentagem («1,1 [11,1]») ou um limite
  («< 1,111», «> 11»).
* «ref_parameter» — o catálogo: id, parâmetro, unidade e o texto de
  referência do laboratório, que muitas vezes tem várias linhas por idade,
  sexo, gravidez ou risco.

Cada parâmetro guarda o texto de referência tal e qual (é o que conta) e um
intervalo numérico tirado dele para o teu caso — adulto, do teu sexo e da tua
idade —, que serve para assinalar ▲/▼. Esse intervalo pode ser corrigido à
mão; uma correção manual não é apagada por uma importação seguinte.
"""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from .db import Database, now_iso

_NUM = r"\d+(?:[.,]\d+)?"
EXCLUDE_MALE = ("mulher", "feminin", "grávid", "gravid", "criança", "crianças", "mese",
                "dias", "semana", "trim", "puberdade", "toxicidade", "deficiência",
                "ferropenia", "insuficiência", "doença", "irc", "hemodiali", "anemia",
                "muito elevado", "elevado:", "moderado:", "risco de diabetes", "diabetes",
                "g2", "g3", "g4", "g5", "aceitável")
EXCLUDE_SECTIONS = ("sexo feminino", "gravidez", "alvos terapêuticos", "deficiência de ferro",
                    "doença inflamatória")
PREFERRED = ("valor de referência", "adulto", "adultos", "normal", "desejável", "baixo",
             "recomendado", "homem", "masculino")


def num(s: str | None) -> float | None:
    if s is None:
        return None
    s = str(s).strip().replace(" ", "")
    if not s:
        return None
    if s.count(",") == 1 and s.count(".") == 0:
        s = s.replace(",", ".")
    elif s.count(",") and s.count("."):
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else \
            s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def parse_value(raw: Any) -> dict:
    """Valor de uma leitura: número, número [percentagem], ou limite."""
    out = {"value": None, "pct": None, "qualifier": None, "text": None}
    if raw is None:
        return out
    if isinstance(raw, (int, float)):
        out["value"] = float(raw)
        return out
    s = str(raw).strip()
    out["text"] = s
    m = re.match(rf"^\s*([<>≤≥]=?)?\s*({_NUM})\s*(?:\[\s*({_NUM})\s*%?\s*\])?\s*$", s)
    if m:
        out["qualifier"] = m.group(1)
        out["value"] = num(m.group(2))
        out["pct"] = num(m.group(3)) if m.group(3) else None
    return out


def _range_in(text: str) -> dict | None:
    """Um intervalo numa porção de texto: «a-b», «< x», «> x», «até x»."""
    t = text.strip().lower()
    m = re.search(rf"({_NUM})\s*(?:-|–|a)\s*({_NUM})", t)
    if m:
        return {"low": num(m.group(1)), "high": num(m.group(2))}
    m = re.search(rf"(?:<|≤|até)\s*=?\s*({_NUM})", t)
    if m:
        return {"low": None, "high": num(m.group(1))}
    m = re.search(rf"(?:>|≥)\s*=?\s*({_NUM})", t)
    if m:
        return {"low": num(m.group(1)), "high": None}
    return None


def _applies(label: str, age: int | None, sex: str) -> tuple[bool, int]:
    """(aplica-se?, especificidade) a uma linha, pelo rótulo."""
    lab = label.lower()
    other = ("mulher", "feminin") if sex == "male" else ("homem", "masculin")
    if any(w in lab for w in other):
        return False, 0
    excl = [w for w in EXCLUDE_MALE if w not in ("mulher", "feminin")]
    if sex == "female":
        excl = [w for w in excl if w not in ("grávid", "gravid")]
    if any(w in lab for w in excl):
        return False, 0
    score = 0
    if any(w in lab for w in PREFERRED):
        score += 1
    if age is not None:
        m = re.search(r"(\d+)\s*a\s*(\d+)\s*anos", lab)
        if m:
            if not int(m.group(1)) <= age <= int(m.group(2)):
                return False, 0
            score += 2
        m = re.search(r"(?:>|mais de)\s*(\d+)\s*anos", lab)
        if m:
            if not age > int(m.group(1)):
                return False, 0
            score += 2
        m = re.search(r"(?:<|até)\s*(\d+)\s*anos", lab)
        if m and not re.search(r"\d+\s*a\s*\d+\s*anos", lab):
            if not age < int(m.group(1)):
                return False, 0
            score += 2
    if re.search(r"homem|masculin", lab) and sex == "male":
        score += 2
    return True, score


def parse_reference(text: str | None, age: int | None = None, sex: str = "male") -> dict:
    """Intervalo para o teu caso, a partir do texto de referência do
    laboratório. Devolve {low, high, ideal_low, ideal_high, source}."""
    out = {"low": None, "high": None, "ideal_low": None, "ideal_high": None, "source": None}
    if not text:
        return out
    best: tuple[int, int, dict, str] | None = None
    section_ok = True
    for i, raw in enumerate(str(text).split("\n")):
        line = raw.strip()
        if not line:
            continue
        low = line.lower()
        has_num = re.search(r"\d", line) is not None
        if line.endswith(":") and not re.search(rf"{_NUM}\s*$", line[:-1]):
            # cabeçalho de secção («Sexo Feminino:», «Gravidez:», «Alvos …:»)
            section_ok = not any(s in low for s in EXCLUDE_SECTIONS)
            if sex == "female" and "sexo feminino" in low:
                section_ok = True
            if sex == "female" and "sexo masculino" in low:
                section_ok = False
            continue
        if not section_ok or not has_num:
            continue
        if ":" in line:
            label, value = line.rsplit(":", 1)
            if not re.search(r"\d", value):          # «5,7 a 6,4: risco de …»
                continue
        else:
            m = re.search(rf"[<>≤≥]?\s*{_NUM}", line)
            label, value = (line[:m.start()], line[m.start():]) if m else (line, "")
            if label.strip().lower().startswith("até"):
                label, value = "", line
        rng = _range_in(value)
        if rng is None:
            continue
        if "ideal" in label.lower():
            out["ideal_low"], out["ideal_high"] = rng["low"], rng["high"]
            continue
        ok, score = _applies(label, age, sex)
        if not ok:
            continue
        cand = (score, -i, rng, line)
        if best is None or cand[:2] > best[:2]:
            best = cand
    if best is not None:
        out["low"], out["high"] = best[2]["low"], best[2]["high"]
        out["source"] = best[3]
    return out


METHOD_WORDS = re.compile(r"lumin|metria|métric|enzim|cinétic|colorim|citometr|imuno|ifcc|icp|"
                          r"potenc|espectro|jaff|hexoquin|quimio|complexo|uricase|urease",
                          re.I)


def short_name(name: str) -> str:
    """Sem o método do laboratório entre parêntesis: «Ferritina
    (Electroquimiluminiscência)» → «Ferritina». Parêntesis que fazem parte do
    nome ficam («HOMA-IR (l. Resistência)»)."""
    name = name.strip()
    for m in re.finditer(r"\s\(", name):
        rest = name[m.end():]
        if METHOD_WORDS.search(rest):
            return name[:m.start()].strip()
    return name


# ------------------------------------------------------------------ esquema


def ensure(db: Database) -> None:
    from . import health
    health.ensure(db)
    db.conn.execute(
        "CREATE TABLE IF NOT EXISTS lab_parameters (id INTEGER PRIMARY KEY,"
        " name TEXT NOT NULL UNIQUE, short_name TEXT, unit TEXT, area TEXT, panel TEXT,"
        " subpanel TEXT, reference_text TEXT, ref_low REAL, ref_high REAL, ideal_low REAL,"
        " ideal_high REAL, ref_source TEXT, ref_manual INTEGER NOT NULL DEFAULT 0,"
        " ext_id INTEGER, sort INTEGER NOT NULL DEFAULT 0)")
    for table, cols in (("lab_reports", (("kind", "TEXT"),)),
                        ("lab_results", (("parameter_id", "INTEGER"), ("area", "TEXT"),
                                         ("panel", "TEXT"), ("subpanel", "TEXT"),
                                         ("value_pct", "REAL"), ("qualifier", "TEXT")))):
        have = {r["name"] for r in db.conn.execute(f"PRAGMA table_info({table})")}
        for c, decl in cols:
            if c not in have:
                db.conn.execute(f"ALTER TABLE {table} ADD COLUMN {c} {decl}")
    db.conn.commit()


def _athlete_age_sex(db: Database) -> tuple[int | None, str]:
    return db.athlete().age, db.get("power_profile_sex", "male")


def parameters(db: Database) -> list[dict]:
    ensure(db)
    return [dict(r) for r in db.conn.execute("SELECT * FROM lab_parameters ORDER BY sort, id")]


def parameter(db: Database, pid: int) -> dict | None:
    ensure(db)
    r = db.conn.execute("SELECT * FROM lab_parameters WHERE id = ?", (pid,)).fetchone()
    return dict(r) if r else None


def parameter_by_name(db: Database, name: str) -> dict | None:
    ensure(db)
    r = db.conn.execute("SELECT * FROM lab_parameters WHERE name = ?", (name.strip(),)).fetchone()
    return dict(r) if r else None


def upsert_parameter(db: Database, name: str, *, unit: str | None = None,
                     reference_text: str | None = None, area: str | None = None,
                     panel: str | None = None, subpanel: str | None = None,
                     ext_id: int | None = None) -> int:
    """Cria ou atualiza um parâmetro. O intervalo calcula-se do texto, a não
    ser que tenha sido corrigido à mão."""
    ensure(db)
    name = name.strip()
    cur = parameter_by_name(db, name)
    age, sex = _athlete_age_sex(db)
    if cur is None:
        sort = db.conn.execute("SELECT COALESCE(MAX(sort), 0) + 1 AS s FROM lab_parameters"
                               ).fetchone()["s"]
        ref = parse_reference(reference_text, age, sex)
        c = db.conn.execute(
            "INSERT INTO lab_parameters(name, short_name, unit, area, panel, subpanel,"
            " reference_text, ref_low, ref_high, ideal_low, ideal_high, ref_source, ext_id, sort)"
            " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (name, short_name(name), unit, area, panel, subpanel, reference_text, ref["low"],
             ref["high"], ref["ideal_low"], ref["ideal_high"], ref["source"], ext_id,
             ext_id if ext_id is not None else sort))
        db.conn.commit()
        return int(c.lastrowid)
    upd = {"unit": unit or cur["unit"], "area": area or cur["area"],
           "panel": panel or cur["panel"], "subpanel": subpanel or cur["subpanel"],
           "reference_text": reference_text if reference_text is not None
           else cur["reference_text"], "ext_id": ext_id if ext_id is not None else cur["ext_id"]}
    sets = ", ".join(f"{k} = ?" for k in upd)
    db.conn.execute(f"UPDATE lab_parameters SET {sets} WHERE id = ?",
                    list(upd.values()) + [cur["id"]])
    if not cur["ref_manual"] and reference_text is not None:
        ref = parse_reference(reference_text, age, sex)
        db.conn.execute("UPDATE lab_parameters SET ref_low=?, ref_high=?, ideal_low=?,"
                        " ideal_high=?, ref_source=? WHERE id = ?",
                        (ref["low"], ref["high"], ref["ideal_low"], ref["ideal_high"],
                         ref["source"], cur["id"]))
    db.conn.commit()
    return cur["id"]


def set_parameter_reference(db: Database, pid: int, low: float | None, high: float | None,
                            unit: str | None = None) -> None:
    """Correção manual: fica, mesmo depois de novas importações."""
    ensure(db)
    db.conn.execute("UPDATE lab_parameters SET ref_low=?, ref_high=?, ref_manual=1,"
                    " unit=COALESCE(?, unit) WHERE id = ?", (low, high, unit, pid))
    db.conn.execute("UPDATE lab_results SET ref_low=?, ref_high=? WHERE parameter_id = ?",
                    (low, high, pid))
    db.conn.commit()


# ------------------------------------------------------------- relatórios


def find_report(db: Database, done_on: str, lab: str, kind: str | None) -> int | None:
    r = db.conn.execute("SELECT id FROM lab_reports WHERE done_on = ? AND COALESCE(lab, '') = ?"
                        " AND COALESCE(kind, '') = ?", (done_on, lab or "", kind or "")).fetchone()
    return r["id"] if r else None


def add_result(db: Database, report_id: int, param_name: str, raw_value: Any, *,
               area: str | None = None, panel: str | None = None,
               subpanel: str | None = None) -> int:
    ensure(db)
    pid = upsert_parameter(db, param_name, area=area, panel=panel, subpanel=subpanel)
    p = parameter(db, pid)
    v = parse_value(raw_value)
    sort = db.conn.execute("SELECT COALESCE(MAX(sort), 0) + 1 AS s FROM lab_results"
                           " WHERE report_id = ?", (report_id,)).fetchone()["s"]
    c = db.conn.execute(
        "INSERT INTO lab_results(report_id, marker, value, value_text, unit, ref_low, ref_high,"
        " comment, sort, parameter_id, area, panel, subpanel, value_pct, qualifier)"
        " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (report_id, p["name"], v["value"], v["text"], p["unit"], p["ref_low"], p["ref_high"],
         None, sort, pid, area, panel, subpanel, v["pct"], v["qualifier"]))
    db.conn.commit()
    return int(c.lastrowid)


def import_xlsx(db: Database, path: str | Path) -> dict:
    """Importa o Excel (folhas «readings» e «ref_parameter»). Voltar a importar
    o mesmo ficheiro atualiza os resultados das mesmas análises (mesma data,
    laboratório e tipo) e mantém as notas e o comentário do médico."""
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError("falta a biblioteca openpyxl (pip install openpyxl)") from e
    ensure(db)
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    names = {n.lower(): n for n in wb.sheetnames}
    report = {"parameters": 0, "reports": 0, "results": 0, "dates": []}
    ref_sheet = names.get("ref_parameter") or next((n for k, n in names.items() if "ref" in k),
                                                   None)
    if ref_sheet:
        rows = list(wb[ref_sheet].iter_rows(values_only=True))
        for row in rows[1:]:
            if not row or row[1] is None:
                continue
            ext = int(row[0]) if isinstance(row[0], (int, float)) else None
            unit = str(row[2]).strip() if len(row) > 2 and row[2] else None
            ref = str(row[3]) if len(row) > 3 and row[3] is not None else None
            upsert_parameter(db, str(row[1]), unit=unit, reference_text=ref, ext_id=ext)
            report["parameters"] += 1
    read_sheet = names.get("readings") or wb.sheetnames[0]
    seen_reports: dict[tuple, int] = {}
    for row in wb[read_sheet].iter_rows(values_only=True):
        if not row or len(row) < 8 or row[6] is None:
            continue
        kind, lab, when, area, panel, subpanel, param, value = row[:8]
        if isinstance(when, datetime):
            day = when.date().isoformat()
        elif isinstance(when, date):
            day = when.isoformat()
        else:
            from .health import _parse_day
            day = _parse_day(when)
            if day is None:
                continue                                    # linha de cabeçalho
        key = (day, str(lab or ""), str(kind or ""))
        if key not in seen_reports:
            rid = find_report(db, *key)
            if rid is None:
                from .health import save_lab_report
                rid = save_lab_report(db, None, done_on=day, lab=key[1])
                db.conn.execute("UPDATE lab_reports SET kind = ? WHERE id = ?", (key[2], rid))
            else:
                db.conn.execute("DELETE FROM lab_results WHERE report_id = ?", (rid,))
            db.conn.commit()
            seen_reports[key] = rid
            report["reports"] += 1
            report["dates"].append(day)
        add_result(db, seen_reports[key], str(param), value, area=_s(area), panel=_s(panel),
                   subpanel=_s(subpanel))
        report["results"] += 1
    wb.close()
    return report


def _s(v: Any) -> str | None:
    return str(v).strip() if v not in (None, "") else None


def grouped_results(db: Database, report_id: int) -> list[tuple[str, str | None, list[dict]]]:
    """[(área, exame, [resultados])] pela ordem do relatório."""
    from .health import lab_results
    out: list[tuple[str, str | None, list[dict]]] = []
    for r in lab_results(db, report_id):
        area = r.get("area") or "—"
        panel = r.get("panel")
        if not out or out[-1][0] != area or out[-1][1] != panel:
            out.append((area, panel, []))
        out[-1][2].append(r)
    return out


def value_text(r: dict) -> str:
    v = r.get("value")
    if v is None:
        return r.get("value_text") or "—"
    s = f"{r['qualifier'] or ''}{v:g}"
    if r.get("value_pct") is not None:
        s += f" [{r['value_pct']:g} %]"
    return s
